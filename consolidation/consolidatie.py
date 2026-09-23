"""Gedragsduur en aantallen consolideren met een noemer per gedragsgroep."""
from pathlib import Path
import math
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

SOURCE_NAME = "10_extra_honden_alle_data_met_modifiers.xlsx"
OVERVIEW_NAME = "10_EXTRA_HONDEN_overzicht data.xlsx"
SCHEMAS = {
    "fases": ["test_id", "fase", "duur_s"],
    "gedrag_groepen": ["gedrag", "groep"],
    "gedragingen": ["test_id", "fase", "gedrag", "duur_s"],
    "out_of_sight": ["test_id", "fase", "groep", "duur_s"],
}


def inspecteer_bron(path):
    """Lees lokale waarden en markeringen; volg geen externe Excel-koppelingen."""
    wb = load_workbook(path, data_only=True, keep_links=False)
    try:
        sheets = pd.DataFrame([
            {"tabblad": ws.title, "rijen": ws.max_row, "kolommen": ws.max_column}
            for ws in wb
        ])
        dogs = pd.DataFrame(wb["Dog data"].values)
        dogs.columns = dogs.iloc[0]
        dogs = dogs.iloc[1:].reset_index(drop=True)[["Test ID", "Dog ID"]]
        dogs.columns = ["test_id", "dog_id"]
        groups = [row[0] for row in wb["Observer groepen"].iter_rows(
            min_row=2, values_only=True) if isinstance(row[0], str)]
        phases = []
        for row in wb["Observer fases data"].iter_rows(min_row=3, values_only=True):
            if row[0] in set(dogs.test_id):
                for phase in range(1, 8):
                    flag = row[phase + 2]
                    if isinstance(flag, str) and flag.strip().lower().startswith("yes"):
                        phases.append({"test_id": row[0], "fase": phase,
                                       "beschikbaarheid": flag})
        yellow = []
        for ws in wb:
            for row in ws:
                for cell in row:
                    color = cell.fill.fgColor
                    if (cell.value is not None and cell.fill.patternType == "solid"
                            and color.type == "rgb" and color.rgb[-6:] == "FFFF00"):
                        yellow.append({"tabblad": ws.title, "cel": cell.coordinate,
                                       "waarde": str(cell.value)})
        return {"tabbladen": sheets, "honden": dogs, "groepen": groups,
                "beschikbare_fases": pd.DataFrame(phases),
                "gele_cellen": pd.DataFrame(yellow)}
    finally:
        wb.close()


def maak_sjabloon(path, bron):
    """Maak alleen een nieuw sjabloon; overschrijf nooit ingevulde invoer."""
    path = Path(path)
    if path.exists():
        raise FileExistsError(f"Sjabloon bestaat al: {path}")
    phases = bron["beschikbare_fases"][["test_id", "fase"]].copy()
    phases["duur_s"] = None
    instructions = [
        "INVOERSJABLOON - ontbrekende duurtijden, geen meetresultaten.",
        "Alle duurtijden zijn numerieke seconden. Vul een expliciete 0 in bij gemeten duur nul.",
        "fases: een rij per aanwezige test/fase. Verwijder alleen werkelijk afwezige fases.",
        "De vooringevulde fases komen uit Observer fases data, voor de 10 tests in Dog data.",
        "gedrag_groepen: koppel elk gedrag aan een groep; meettype is duur of aantal.",
        "gedragingen: een rij per test/fase/gedrag. Vul duur_s OF aantal in volgens meettype.",
        "Laat de andere meetkolom leeg. Vul gemeten nullen expliciet in.",
        "Aantal / zichtbare seconden geeft frequentie per seconde, geen percentage.",
        "out_of_sight: een rij per test/fase/groep, ook bij duur 0.",
        "Sommeer losse gebeurtenissen eerst per fase; dubbele sleutels worden geweigerd.",
        "Overlappende out-of-sight-intervallen binnen een groep niet dubbel tellen.",
        "Gebruik uitsluitend fases 1 t/m 7. De andere fases worden niet geconsolideerd.",
        "Er zijn geen gedragstijden of out-of-sight-tijden gevonden in het overzichtsbestand.",
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame({"uitleg": instructions}).to_excel(writer, sheet_name="LEESMIJ", index=False)
        for name, columns in SCHEMAS.items():
            frame = phases if name == "fases" else pd.DataFrame(columns=columns)
            if name == "gedrag_groepen":
                frame["meettype"] = pd.Series(dtype="str")
            if name == "gedragingen":
                frame["aantal"] = pd.Series(dtype="float")
            frame.to_excel(writer, sheet_name=name, index=False)
        pd.DataFrame({"groep": bron["groepen"]}).to_excel(
            writer, sheet_name="groepen_referentie", index=False)
        for ws in writer.book:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = 28
        writer.book["LEESMIJ"].column_dimensions["A"].width = 115
        writer.book["groepen_referentie"].column_dimensions["A"].width = 65
        for cell in writer.book["out_of_sight"][1]:
            cell.fill = PatternFill("solid", fgColor="FFFF00")
    return path


def lees_invoer(path):
    with pd.ExcelFile(path, engine="openpyxl") as book:
        missing = set(SCHEMAS) - set(book.sheet_names)
        if missing:
            raise ValueError(f"Ontbrekende tabbladen: {sorted(missing)}")
        return {name: pd.read_excel(book, sheet_name=name) for name in SCHEMAS}


def _controleer(frame, name, keys, numeric=True):
    missing = set(SCHEMAS[name]) - set(frame.columns)
    if missing:
        raise ValueError(f"{name}: ontbrekende kolommen {sorted(missing)}")
    optional = {"gedrag_groepen": ["meettype"], "gedragingen": ["aantal"]}
    columns = SCHEMAS[name] + [c for c in optional.get(name, []) if c in frame]
    frame = frame[columns].copy()
    if frame.empty:
        raise ValueError(f"{name}: geen invoer.")
    for column in keys:
        if frame[column].isna().any():
            raise ValueError(f"{name}: ontbrekende waarde in {column}.")
        if column != "fase":
            if not frame[column].map(lambda x: isinstance(x, str) and bool(x.strip())).all():
                raise ValueError(f"{name}: ongeldige tekst in {column}.")
            frame[column] = frame[column].str.strip()
    if "fase" in frame:
        valid = frame.fase.map(lambda x: not isinstance(x, bool)
                               and isinstance(x, (int, float))
                               and math.isfinite(x) and x == int(x) and 1 <= x <= 15)
        if not valid.all():
            raise ValueError(f"{name}: ongeldige fase (verwacht geheel getal 1..15).")
        frame["fase"] = frame.fase.astype(int)
    if frame.duplicated(keys).any():
        raise ValueError(f"{name}: dubbele sleutel {keys}.")
    if numeric:
        valid = frame.duur_s.map(lambda x: not isinstance(x, bool)
                                and isinstance(x, (int, float))
                                and math.isfinite(x) and x >= 0)
        if not valid.all():
            raise ValueError(f"{name}: duur_s moet ingevulde, eindige, niet-negatieve seconden bevatten.")
        frame["duur_s"] = frame.duur_s.astype(float)
    return frame


def consolideer(fases, gedrag_groepen, gedragingen, out_of_sight, tolerantie_s=1e-9,
               geselecteerde_fases=tuple(range(1, 8))):
    """Som gedrag / (som fase - som groeps-OOS), uitsluitend fases 1..7.

    Invoer bevat fasesommen, geen losse gebeurtenissen. Iedere combinatie van
    aanwezige fase en gedefinieerd gedrag/groep moet expliciet aanwezig zijn.
    Een volledig onzichtbare groep krijgt een lege fractie met een duidelijke status.
    """
    gekozen = tuple(geselecteerde_fases)
    if (not gekozen or len(set(gekozen)) != len(gekozen)
            or any(type(x) is not int or not 1 <= x <= 15 for x in gekozen)):
        raise ValueError("Ongeldige selectie van fases: gebruik unieke gehele fasenummers 1..15.")
    f = _controleer(fases, "fases", ["test_id", "fase"])
    m = _controleer(gedrag_groepen, "gedrag_groepen", ["gedrag"], numeric=False)
    if m.groep.isna().any() or not m.groep.map(
            lambda x: isinstance(x, str) and bool(x.strip())).all():
        raise ValueError("gedrag_groepen: ontbrekende of ongeldige groep.")
    m["groep"] = m.groep.str.strip()
    # Oude invoer met alleen duur_s blijft bruikbaar. Een aanwezige maar lege
    # meettype-kolom wordt bewust niet automatisch als duur geinterpreteerd.
    if "meettype" not in m:
        m["meettype"] = "duur"
    if not m.meettype.isin(["duur", "aantal"]).all():
        raise ValueError("meettype moet expliciet duur of aantal zijn.")
    b = _controleer(gedragingen, "gedragingen", ["test_id", "fase", "gedrag"], numeric=False)
    if not set(b.gedrag) <= set(m.gedrag):
        raise ValueError("Onbekend gedrag: ontbrekende groepskoppeling.")
    if "aantal" not in b:
        b["aantal"] = float("nan")
    b = b.merge(m, on="gedrag", validate="many_to_one")
    for column, kind in [("duur_s", "duur"), ("aantal", "aantal")]:
        selected = b.meettype == kind
        valid = b.loc[selected, column].map(
            lambda x: not isinstance(x, bool) and isinstance(x, (int, float))
            and math.isfinite(x) and x >= 0
            and (kind != "aantal" or x == int(x)))
        if not valid.all():
            raise ValueError(f"{column}: ontbrekende of ongeldige meting; aantallen moeten geheel zijn.")
        if b.loc[~selected, column].notna().any():
            raise ValueError(f"{column}: laat de niet-toepasselijke meetkolom leeg.")
        b[column] = b[column].astype(float)
    o = _controleer(out_of_sight, "out_of_sight", ["test_id", "fase", "groep"])
    f = f[f.fase.isin(gekozen)].copy()
    b = b[b.fase.isin(gekozen)].copy()
    o = o[o.fase.isin(gekozen)].copy()
    if f.empty:
        raise ValueError(f"Geen aanwezige fases binnen de selectie {gekozen}.")
    if not set(b.gedrag) <= set(m.gedrag):
        raise ValueError("Onbekend gedrag: ontbrekende groepskoppeling.")
    if not set(o.groep) <= set(m.groep):
        raise ValueError("Onbekende out-of-sight-groep.")
    phase_keys = ["test_id", "fase"]
    for frame, name in [(b, "gedragingen"), (o, "out_of_sight")]:
        joined = frame.merge(f[phase_keys], on=phase_keys, how="left", indicator=True)
        if (joined._merge != "both").any():
            raise ValueError(f"{name}: test/fase ontbreekt in fases.")
    expected_b = f[phase_keys].merge(m[["gedrag"]], how="cross")
    expected_o = f[phase_keys].merge(m[["groep"]].drop_duplicates(), how="cross")
    for expected, actual, keys, name in [
        (expected_b, b, phase_keys + ["gedrag"], "gedragingen"),
        (expected_o, o, phase_keys + ["groep"], "out_of_sight"),
    ]:
        check = expected.merge(actual[keys], on=keys, how="left", indicator=True)
        if (check._merge != "both").any():
            raise ValueError(f"{name}: ontbrekende metingen; vul gemeten nullen expliciet in.")
    visibility = expected_o.merge(
        f.rename(columns={"duur_s": "fase_duur_s"}), on=phase_keys,
        validate="many_to_one").merge(
        o.rename(columns={"duur_s": "out_of_sight_s"}),
        on=phase_keys + ["groep"], validate="one_to_one")
    visibility["zichtbaar_s"] = visibility.fase_duur_s - visibility.out_of_sight_s
    if (visibility.zichtbaar_s < -tolerantie_s).any():
        raise ValueError("Out of sight is groter dan de faseduur.")
    visibility["zichtbaar_s"] = visibility.zichtbaar_s.clip(lower=0)
    detail = b.rename(columns={"duur_s": "gedrag_s"}).merge(
        visibility, on=phase_keys + ["groep"], validate="many_to_one")
    if (detail.gedrag_s > detail.zichtbaar_s + tolerantie_s).any():
        raise ValueError("Gedragsduur is groter dan de zichtbare tijd binnen een fase.")
    if ((detail.zichtbaar_s == 0) & (detail.aantal > 0)).any():
        raise ValueError("Positief aantal zonder zichtbare tijd binnen een fase.")
    sums = detail.groupby(["test_id", "groep", "gedrag", "meettype"], as_index=False).agg(
        gedrag_s=("gedrag_s", lambda x: x.sum(min_count=1)),
        aantal=("aantal", lambda x: x.sum(min_count=1)),
        totale_faseduur_s=("fase_duur_s", "sum"),
        out_of_sight_s=("out_of_sight_s", "sum"),
        aantal_fases=("fase", "nunique"),
        aanwezige_fases=("fase", lambda x: ", ".join(map(str, sorted(x)))),
    )
    # Eerst sommeren en dan delen, geen gemiddelde van fasepercentages.
    sums["zichtbaar_s"] = (sums.totale_faseduur_s - sums.out_of_sight_s).clip(lower=0)
    sums["fractie"] = sums.gedrag_s / sums.zichtbaar_s.where(sums.zichtbaar_s > 0)
    sums["percentage"] = 100 * sums.fractie
    sums["frequentie_per_s"] = sums.aantal / sums.zichtbaar_s.where(sums.zichtbaar_s > 0)
    sums["frequentie_per_min"] = 60 * sums.frequentie_per_s
    sums["status"] = sums.zichtbaar_s.map(
        lambda x: "ok" if x > 0 else "geen_zichtbare_tijd")
    return sums, detail.sort_values(phase_keys + ["groep", "gedrag"]).reset_index(drop=True)


def exporteer(path, resultaat, detail):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        resultaat.to_excel(writer, sheet_name="consolidatie", index=False)
        detail.to_excel(writer, sheet_name="controle_per_fase", index=False)
        for ws in writer.book:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = 24


def demo_invoer():
    """Volledig fictief voorbeeld, nooit data van de echte honden."""
    return {
        "fases": pd.DataFrame([["DEMO", 1, 100], ["DEMO", 2, 200]],
                              columns=SCHEMAS["fases"]),
        "gedrag_groepen": pd.DataFrame([["zitten", "houding"], ["kijken", "aandacht"]],
                                       columns=SCHEMAS["gedrag_groepen"]),
        "gedragingen": pd.DataFrame([
            ["DEMO", 1, "zitten", 20], ["DEMO", 2, "zitten", 10],
            ["DEMO", 1, "kijken", 10], ["DEMO", 2, "kijken", 20]],
            columns=SCHEMAS["gedragingen"]),
        "out_of_sight": pd.DataFrame([
            ["DEMO", 1, "houding", 10], ["DEMO", 2, "houding", 20],
            ["DEMO", 1, "aandacht", 40], ["DEMO", 2, "aandacht", 20]],
            columns=SCHEMAS["out_of_sight"]),
    }
