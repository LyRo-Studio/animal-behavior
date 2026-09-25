"""Import van het ruwe Results-blad van de Observer-export; behoud iedere modifiercombinatie."""
from pathlib import Path
import hashlib
import math
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from consolidation import ethogram_definition
from consolidation.ethogram_definition import normalise
from consolidation.consolidatie import consolideer, fase_label

# Welke Behaviour group en Out of Sight bij een kolom horen, komt uit de
# definitie die uit het ethogram gegenereerd is (ticket #150), nooit uit de
# kolompositie in de export.
DEFINITIE = ethogram_definition.load()
# Het enige blad dat gelezen wordt (ticket #151): de export zoals Observer
# hem maakt. Fase, Geslacht hond en Observers containerkolommen worden nooit
# gelezen.
BLAD = "Results"
EIGENAAR_AANWEZIG = "Aanwezigheid FP in fase"
VERPLICHTE_KOLOMMEN = ["Observations", "Test ID", "Dog ID", EIGENAAR_AANWEZIG, "Duration"]
TOLERANTIE_S = 0.001  # Observer exporteert seconden met een beperkte decimale precisie.
ME, ZE = tuple(range(1, 8)), tuple(range(9, 16))
# De vier Consolidation levels (ticket #152), elk met een eigen resultaatblad.
# Vertrek eigenaar (fase 8) telt in geen enkel niveau.
NIVEAUS = {
    "per_phase": {"fases": ME + ZE, "per_fase": True},
    "with_owner": {"fases": ME},
    "without_owner": {"fases": ZE},
    "combined": {"fases": ME + ZE},
}
GEBRUIKTE_FASES = set(ME + ZE)


def fase_uit_observatie(value):
    match = re.fullmatch(r"(T\d+)_.+_F(\d+)", str(value))
    if not match:
        raise ValueError(f"Malformed Observation name {value!r}: expected <Test ID>_..._F<phase>.")
    phase = int(match[2])
    if not 1 <= phase <= 15:
        raise ValueError(f"Phase outside 1-15 in Observation name {value!r}.")
    return match[1], phase


def _bekende_gedragingen(definitie):
    """Iedere ethogramnaam en exportalias (genormaliseerd), langste eerst."""
    per_naam = {b["name"]: b for b in definitie["behaviours"]}
    namen = {normalise(b["name"]): b for b in definitie["behaviours"]}
    namen.update({normalise(alias): per_naam[doel]
                  for alias, doel in definitie["behaviour_aliases"].items()})
    return sorted(namen.items(), key=lambda item: len(item[0]), reverse=True)


def splits_kop(header, bekend):
    """'Total duration Tail tucked <No Modifier>' -> ('duur', gedrag, '<No Modifier>').

    Het gedrag is de langste bekende naam waarmee de kop begint, gevolgd door
    de modifier. None als geen enkele bekende naam past.
    """
    kind = "duur" if header.startswith("Total duration ") else "aantal"
    rest = ethogram_definition.tidy(header.removeprefix("Total duration ").removeprefix("Total number "))
    for naam, gedrag in bekend:
        if rest.casefold().startswith(naam + " "):
            return kind, gedrag, rest[len(naam) + 1:]
    return None


def modifier_toegestaan(modifier, gedrag, definitie):
    """Past iedere modifierwaarde in een eigen, door het ethogram toegestane categorie?"""
    if modifier == "<No Modifier>":
        return True
    categorieen = [{normalise(v) for v in definitie["modifier_categories"][c]}
                   for c in gedrag["modifier_categories"]]
    aliassen = {normalise(k): normalise(v) for k, v in definitie["modifier_aliases"].items()}
    waarden = [aliassen.get(normalise(w), normalise(w)) for w in modifier.split(";")]
    kandidaten = [{i for i, c in enumerate(categorieen) if w in c} for w in waarden]
    if len(waarden) > len(categorieen) or not all(kandidaten):
        return False
    return len(waarden) < 2 or any(a != b for a in kandidaten[0] for b in kandidaten[1])


def _deel_kolommen_in(headers, source, definitie):
    """Deel iedere Total duration/number-kolom in volgens de definitie.

    Een kolom zonder bekend gedrag faalt. De Out of Sight van een groep is
    alleen verplicht wanneer minstens een gedrag van die groep geexporteerd is.
    `ongelezen` zijn de `Total number Out of sight …`-kolommen: de enige
    Total-kolommen waarvan de cellen nooit gelezen (en dus nooit gecontroleerd)
    worden.
    """
    bekend = _bekende_gedragingen(definitie)
    groepen = {g["name"]: g for g in definitie["groups"]}
    mapping, excluded, notes = [], [], []
    oos_kolommen, oos_kandidaten, event_duur, per_gedrag = {}, [], [], {}
    niet_ondersteund = {}
    for header, c in headers.items():
        if not header.startswith(("Total duration ", "Total number ")):
            continue
        letter = source.cell(1, c).column_letter
        gesplitst = splits_kop(header, bekend)
        if gesplitst is None:
            raise ValueError(
                f"Unknown behaviour in column {letter} ({header}): it is not in the Ethogram "
                "definition. Add it to the Ethogram and regenerate the definition.")
        kind, gedrag, modifier = gesplitst
        groep = groepen[gedrag["group"]]
        per_gedrag.setdefault(gedrag["name"], []).append(c)
        rij = {"bronkolom": letter, "bronkop": header}
        if gedrag["name"] == groep["out_of_sight"]:
            if kind == "duur" and modifier == "<No Modifier>":
                oos_kolommen[groep["name"]] = c
            oos_kandidaten.append((rij, groep["name"], c, kind))
        elif groep["excluded"] or gedrag["name"] in definitie["excluded_behaviours"]:
            excluded.append({**rij, "reden": groep["excluded"]
                             or definitie["excluded_behaviours"][gedrag["name"]]})
        elif groep["scored_phases"] != ethogram_definition.ALL_PHASES:
            niet_ondersteund.setdefault(groep["name"], groep["scored_phases"])
            excluded.append({**rij, "reden": "Not yet supported: this group is only scored in "
                             "some phases (Scoring plan)."})
        elif gedrag["kind"] == "Event" and kind == "duur":
            event_duur.append((header, c))
            excluded.append({**rij, "reden": "Event duration: not a meaningful measure; "
                             "only the Event's frequency is reported."})
        else:
            if not modifier_toegestaan(modifier, gedrag, definitie):
                notes.append(_melding("unknown_modifier", header,
                                      f"The Ethogram doesn't allow modifier {modifier!r} for "
                                      f"{gedrag['name']}; its values are kept."))
            mapping.append({"gedrag": header, "basisgedrag": gedrag["name"],
                            "modifier": modifier, "groep": groep["name"], "meettype": kind,
                            "code": gedrag["code"], "state_event": gedrag["kind"],
                            "bronkolom": letter, "kolomnummer": c})
    if not mapping:
        raise ValueError("The export has no dog behaviour columns to consolidate.")
    gebruikt = {m["groep"] for m in mapping}
    for naam in sorted(gebruikt):
        oos = groepen[naam]["out_of_sight"]
        if oos and naam not in oos_kolommen:
            raise ValueError(
                f"Missing Out of Sight column for {naam}: the export has behaviours of this "
                f"group but no 'Total duration {oos} <No Modifier>' column.")
    ongelezen = set()
    for rij, groep, c, kind in oos_kandidaten:
        if kind == "aantal":
            reden = "Out of Sight count: never read."
            ongelezen.add(c)
        elif oos_kolommen.get(groep) != c:
            reden = "Out of Sight modifier column: never used."
        elif groep in gebruikt:
            reden = "Out of Sight correction column: its duration corrects the group's denominator."
        else:
            reden = "Out of Sight column ignored: none of its group's behaviours were exported."
        excluded.append({**rij, "reden": reden})
    for naam, fases in niet_ondersteund.items():
        notes.append(_melding("not_yet_supported", naam,
                              f"{naam} is only scored in phases {fases} of each Condition. "
                              "Its columns are excluded until the Scoring plan is supported."))
    oos_per_groep = {naam: oos_kolommen.get(naam) for naam in gebruikt}
    for m in mapping:
        oc = oos_per_groep[m["groep"]]
        m["oos_kolom"] = source.cell(1, oc).column_letter if oc else "geen"
    return mapping, excluded, notes, oos_per_groep, event_duur, per_gedrag, ongelezen


def _beschikbaarheid(per_gedrag, niet_nul, definitie):
    """Ieder hondgedrag uit het ethogram: niet, met alleen nullen of met waarden geexporteerd."""
    groepen = {g["name"]: g for g in definitie["groups"]}
    rijen = []
    for gedrag in definitie["behaviours"]:
        if groepen[gedrag["group"]]["excluded"]:
            continue
        kolommen = per_gedrag.get(gedrag["name"], [])
        status = ("not_exported" if not kolommen else
                  "exported_nonzero" if niet_nul & set(kolommen) else "exported_all_zero")
        rijen.append({"behaviour": gedrag["name"], "code": gedrag["code"],
                      "group": gedrag["group"], "status": status,
                      "exported_columns": len(kolommen)})
    return pd.DataFrame(rijen)


def _melding(soort, variabele, melding, test_id="", fase=""):
    return {"type": soort, "test_id": test_id, "fase": fase, "variabele": variabele,
            "melding": melding}


def _is_meetkolom(header):
    return header == "Duration" or header.startswith(("Total duration ", "Total number "))


def _plaats(cell, header):
    return f"{BLAD}!{cell.coordinate} ({header})"


def _ingevuld(cell, header):
    """De waarde van een verplichte cel; een blanke cel faalt en noemt de cel."""
    value = cell.value
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"Blank cell {_plaats(cell, header)}")
    return value


def _meetwaarde(cell, header):
    """Een meetcel als seconden of aantal. Observers '-' is een gemeten nul.

    Een blanke of niet-numerieke cel faalt en noemt de cel.
    """
    value = cell.value
    if isinstance(value, str) and value.strip() == "-":
        return 0.0
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"Blank cell {_plaats(cell, header)}: every exported cell must hold "
                         "a number, or '-' for a measured zero.")
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Not a number in cell {_plaats(cell, header)}: {str(value)[:40]!r}.")
    if not math.isfinite(value) or value < 0:
        raise ValueError(f"Invalid value in cell {_plaats(cell, header)}: {value!r}; "
                         "measurements are finite and never negative.")
    return float(value)


def _kolommen(blad):
    """{kop: kolomnummer} van alle kolommen die gelezen worden.

    Een dubbele verplichte of Total-kolom faalt: welke van de twee telt, is
    niet te zeggen. Andere kolommen worden nooit gelezen, dus hun koppen doen
    er niet toe.
    """
    per_kop = {}
    for cell in blad[1]:
        header = "" if cell.value is None else str(cell.value)
        if header in VERPLICHTE_KOLOMMEN or _is_meetkolom(header):
            per_kop.setdefault(header, []).append(cell)
    for header, cells in per_kop.items():
        if len(cells) > 1:
            letters = ", ".join(c.column_letter for c in cells)
            raise ValueError(f"Duplicate column in {BLAD}: {header} (columns {letters})")
    for header in VERPLICHTE_KOLOMMEN:
        if header not in per_kop:
            raise ValueError(f"Missing column in {BLAD}: {header}")
    return {header: cell.column for header, (cell,) in per_kop.items()}


def lees_observer(path):
    """Lees iedere fase van ieder niveau uit het ruwe Results-blad van de Observer-export.

    Alleen dat blad wordt gelezen (ticket #151); handgemaakte werkkopieen zijn
    niet nodig. '-' is een gemeten nul; een blanke of niet-numerieke meetcel
    faalt. Fase wordt uit Observations gehaald: de bronkolom Fase nummert
    deel 2 opnieuw van 1 tot 7 en wordt nooit gelezen.
    """
    try:
        wb = load_workbook(path, data_only=False, keep_links=False)
    except KeyError as exc:
        # openpyxl meldt sommige ontbrekende OOXML-onderdelen als KeyError.
        raise ValueError("The uploaded file is not a valid Excel workbook.") from exc
    try:
        if BLAD not in wb.sheetnames:
            raise ValueError(f"Missing sheet: {BLAD}")
        source = wb[BLAD]
        headers = _kolommen(source)
        (mapping, excluded, notes, oos_columns, event_duur, per_gedrag,
         ongelezen) = _deel_kolommen_in(headers, source, DEFINITIE)
        # Duration en iedere Total-kolom, behalve de Out of Sight-aantallen.
        te_lezen = {header: c for header, c in headers.items()
                    if _is_meetkolom(header) and c not in ongelezen}
        definitions = pd.DataFrame(mapping)
        phases, behaviors, invisibility, selection = [], [], [], []
        niet_nul = set()
        for row in range(2, source.max_row + 1):
            if all(cell.value is None for cell in source[row]):
                continue  # Een lege rij is geen observatie.
            observation, test, owner_flag = (
                _ingevuld(source.cell(row, headers[h]), h)
                for h in ("Observations", "Test ID", EIGENAAR_AANWEZIG))
            label_test, phase = fase_uit_observatie(observation)
            dog = source.cell(row, headers["Dog ID"]).value
            owner_flag = str(owner_flag).lower()
            if owner_flag not in {"true", "false"} or (owner_flag == "true") != (phase <= 7):
                raise ValueError(
                    f"The owner-present flag ({EIGENAAR_AANWEZIG}) contradicts the phase of "
                    f"Observation {observation}: it must be True exactly for phases 1-7.")
            if test != label_test:
                notes.append(_melding("observation_name", observation,
                                      f"The Observation name's Test ID ({label_test}) differs "
                                      f"from the Test ID column; {test} is used.", test, phase))
            selection.append({"observatie": observation, "test_id": test, "dog_id": dog,
                              "fase": phase, "fase_in_deel": phase if phase <= 7 else phase - 8,
                              "deel": 1 if phase <= 7 else 2, "met_eigenaar": owner_flag == "true",
                              "opgenomen": phase in GEBRUIKTE_FASES, "bronrij": row})
            if phase not in GEBRUIKTE_FASES:
                continue
            # Controleer iedere gelezen kolom, ook uitgesloten protocolkolommen.
            waarden = {}
            for header, c in te_lezen.items():
                waarden[c] = _meetwaarde(source.cell(row, c), header)
                if waarden[c] > 0:
                    niet_nul.add(c)
            phases.append({"test_id": test, "fase": phase, "duur_s": waarden[headers["Duration"]]})
            for group, oc in oos_columns.items():
                invisibility.append({"test_id": test, "fase": phase, "groep": group,
                                     "duur_s": waarden[oc] if oc else 0.0})
            for m in mapping:
                value = waarden[m["kolomnummer"]]
                behaviors.append({"test_id": test, "fase": phase, "gedrag": m["gedrag"],
                                  "duur_s": value if m["meettype"] == "duur" else float("nan"),
                                  "aantal": value if m["meettype"] == "aantal" else float("nan")})
        for header, c in event_duur:
            if c in niet_nul:
                notes.append(_melding("event_duration", header,
                                      "This Event has a nonzero duration; Observer's coding may "
                                      "have changed. Only its frequency is reported."))
        selected = pd.DataFrame(selection)
        if selected.empty:
            raise ValueError(f"The {BLAD} sheet has no Observations.")
        if not phases:
            raise ValueError(f"The {BLAD} sheet has no Observations in Observer phases 1-7 or 9-15.")
        dubbel = selected[selected.duplicated(["test_id", "fase"])]
        if not dubbel.empty:
            raise ValueError(f"Duplicate test and phase: {dubbel.test_id.iloc[0]} phase "
                             f"{dubbel.fase.iloc[0]} appears more than once.")
        dogs = selected[["test_id", "dog_id"]].drop_duplicates()
        fout = dogs[dogs.test_id.duplicated(keep=False) | dogs.dog_id.isna()].test_id.unique()
        if len(fout):
            raise ValueError(f"Missing or conflicting Dog ID for test {', '.join(map(str, fout))}.")
        return {"invoer": {"fases": pd.DataFrame(phases),
                           "gedrag_groepen": definitions[["gedrag", "groep", "meettype"]],
                           "gedragingen": pd.DataFrame(behaviors),
                           "out_of_sight": pd.DataFrame(invisibility)},
                "definities": definitions, "honden": dogs, "selectie": selected,
                "uitgesloten": pd.DataFrame(excluded), "meldingen": pd.DataFrame(notes),
                "beschikbaarheid": _beschikbaarheid(per_gedrag, niet_nul, DEFINITIE)}
    finally:
        wb.close()


def _breed(resultaat, sleutel, variabelen):
    """Een rij per `sleutel`; duurkolommen bevatten fracties en aantal-kolommen
    frequenties per seconde, in de volgorde van `variabelen`."""
    if resultaat.empty:
        return pd.DataFrame(columns=[*sleutel, *variabelen])
    wide = resultaat.pivot(index=sleutel, columns="gedrag", values="geconsolideerde_waarde")
    wide = wide.reindex(columns=variabelen).reset_index()
    wide.columns.name = None
    return wide


def bereken_observer(bron):
    """Ieder niveau (NIVEAUS) uit een gelezen export: een breed resultaatblad
    per niveau, plus de lange tabellen en meldingen."""
    resultaat, detail, noemers = consolideer(**bron["invoer"], niveaus=NIVEAUS,
                                            tolerantie_s=TOLERANTIE_S)
    labels = bron["definities"][["gedrag", "basisgedrag", "modifier", "bronkolom", "oos_kolom"]]
    resultaat = resultaat.merge(labels, on="gedrag", validate="many_to_one").merge(
        bron["honden"], on="test_id", validate="many_to_one")
    resultaat["geconsolideerde_waarde"] = resultaat.fractie.where(
        resultaat.meettype == "duur", resultaat.frequentie_per_s)
    notes = bron["meldingen"].to_dict("records")
    for row in detail[detail.gedrag_s > detail.zichtbaar_s].itertuples():
        notes.append(_melding("rounding", row.gedrag,
                              f"Behaviour duration {row.gedrag_s} s against visible time "
                              f"{row.zichtbaar_s} s: within the {TOLERANTIE_S} s export tolerance; "
                              "values kept unchanged.", row.test_id, row.fase))
    groepen = detail.drop_duplicates(["test_id", "fase", "groep"])
    for row in groepen[groepen.out_of_sight_s > groepen.fase_duur_s].itertuples():
        notes.append(_melding("rounding", row.groep,
                              f"Out of Sight {row.out_of_sight_s} s against Duration "
                              f"{row.fase_duur_s} s: within the {TOLERANTIE_S} s export tolerance, "
                              "so no visible time.", row.test_id, row.fase))
    variabelen = list(bron["definities"].gedrag)
    aanwezig = bron["invoer"]["fases"].groupby("test_id").fase.agg(set)
    per_niveau = {}
    for niveau, config in NIVEAUS.items():
        van_niveau = resultaat[resultaat.level == niveau]
        if config.get("per_fase"):
            wide = _breed(van_niveau, ["test_id", "dog_id", "fase"], variabelen).rename(
                columns={"fase": "observer_phase"})
            labels_fase = wide.observer_phase.map(fase_label)
            wide.insert(3, "condition", labels_fase.str[:2])
            wide.insert(4, "phase", labels_fase.str[3:])
            per_niveau[niveau] = wide
            continue
        rijen = []
        for test in bron["honden"].itertuples():
            gebruikt = aanwezig[test.test_id] & set(config["fases"])
            if not gebruikt:
                notes.append(_melding("no_phases", niveau,
                                      f"{test.test_id} has none of the {niveau} phases, so it has "
                                      f"no row on {niveau}.", test.test_id))
                continue
            ontbrekend = set(config["fases"]) - gebruikt
            rijen.append({"test_id": test.test_id, "dog_id": test.dog_id,
                          "phases_used": ", ".join(map(fase_label, sorted(gebruikt))),
                          "missing_phases": ", ".join(map(fase_label, sorted(ontbrekend))) or None,
                          "status": "incomplete" if ontbrekend else "ok"})
        kop = pd.DataFrame(rijen, columns=["test_id", "dog_id", "phases_used",
                                            "missing_phases", "status"])
        per_niveau[niveau] = kop.merge(_breed(van_niveau, ["test_id"], variabelen), on="test_id",
                                       how="left", validate="one_to_one")
    noemers = noemers.merge(bron["honden"], on="test_id", validate="many_to_one")
    return {"per_niveau": per_niveau, "details": resultaat, "noemers": noemers,
            "meldingen": pd.DataFrame(notes)}


def schrijf_resultaat(path, bron, berekend, source_path):
    """Een Engelstalig werkboek met ieder niveau (NIVEAUS) uit een upload.

    `bron` en `berekend` zijn de uitkomst van lees_observer en bereken_observer.
    """
    path = Path(path)
    path.parent.mkdir(exist_ok=True)
    uitleg = [
        "CONSOLIDATION RESULT - every level from one Observer export.",
        "per_phase: every Observer phase 1-7 and 9-15 on its own; one row per test and phase, with observer_phase, "
        "condition (ME/ZE) and phase (F1-F7).",
        "with_owner: Observer phases 1-7 (ME, the owner is present).",
        "without_owner: Observer phases 9-15 (ZE, the owner is absent).",
        "combined: Observer phases 1-7 and 9-15 together. Owner departure (phase 8) never counts.",
        "Phases come from the Observation name; the owner-present flag is checked against them.",
        "with_owner, without_owner, combined: one row per test that has at least one of the level's phases, with "
        "phases_used, missing_phases and status: ok (every phase present) or incomplete (calculated from the phases "
        "present; the missing ones are listed). A test with none of a level's phases has no row there, with a warning.",
        "Every original behaviour and modifier column is kept separately, on every result sheet.",
        "Total duration columns: sum of behaviour duration / (sum of phase Duration - sum of the group's Out of Sight duration). Fractions, not percentages.",
        "Total number columns: sum of counts / the same denominator. Unit: occurrences per visible second.",
        "Sums first, then one division: never an average of phase values. combined is never an average of with_owner and without_owner.",
        "Modifier combinations are never summed into a base behaviour; each is only summed over phases.",
        "A group without an Out of Sight behaviour (Vocalisation) uses the full phase Duration.",
        "Out of Sight counts are not time: only Out of Sight durations are subtracted.",
        "Behaviour groups and Out of Sight corrections come from the Ethogram definition, never from column positions. "
        "Stiffening up belongs to exploration/self-maintenance.",
        "A behaviour Observer didn't export has no result column; an exported column holding 0 is a measured zero.",
        "Events report only a frequency; their duration columns are listed in excluded.",
        "TP/FP protocol behaviours and First contact with TP (a separate F1 measure) are recognised by name and listed in excluded.",
        "Distance to TP/FP, Location and Dog following are only scored in some phases: until the Scoring plan is supported, "
        "their columns are listed in excluded, with a warning.",
        f"No visible time: visible time of {TOLERANTIE_S} s or less, per phase or summed over a level, gives a blank "
        "value with status no_visible_time (in details and denominators); never 0 and never a division.",
        "Out of Sight above Duration by more than the tolerance fails the upload; within it, visible time is 0 and "
        "there is a rounding warning.",
        "Only the export's raw Results sheet is read, exactly as Observer produces it. The source file is not changed.",
        "Observer's '-' is a measured 0. A blank or non-numeric cell in an exported column fails the upload.",
        "Never read: Fase, Geslacht hond, Observer's container columns, Total number Out of sight columns "
        "and the owner-departure (phase 8) rows. Every other exported cell is checked, used or not.",
        f"Maximum tolerance on individual times: {TOLERANTIE_S} s for export rounding; deviations are listed in warnings.",
        "details: every numerator, denominator, percentage, frequency per second and per minute, and status, per level "
        "(fase is the Observer phase, on per_phase rows only).",
        "denominators: Duration, Out of Sight and visible time per level, test, phase (per_phase only) and group, "
        "with the phases used and the status, so every blank can be explained.",
        "phase_selection: every source row and the levels it was used in; variables: every result column; excluded: columns left out, with the reason.",
        "availability: every dog behaviour in the Ethogram, as not_exported (absent from the export, so absent from the results), "
        "exported_all_zero (a measured zero) or exported_nonzero.",
        "Diagnostic sheets keep the calculation's own column names: fase = Observer phase number, groep = Behaviour group, "
        "gedrag = result column, basisgedrag = behaviour, duur = duration, aantal = count, zichtbaar = visible time, "
        "fractie = fraction, frequentie = frequency.",
        f"Source: {Path(source_path).name}",
        f"SHA256 source: {hashlib.sha256(Path(source_path).read_bytes()).hexdigest()}",
    ]

    selectie = bron["selectie"].drop(columns="opgenomen")
    selectie["levels"] = [", ".join(n for n, c in NIVEAUS.items() if fase in c["fases"]) or None
                          for fase in selectie.fase]
    details = berekend["details"]
    details = details[["level", "test_id", "fase",
                       *(c for c in details.columns if c not in {"level", "test_id", "fase"})]]
    noemers = berekend["noemers"].rename(columns={
        "fase": "observer_phase", "groep": "group", "totale_faseduur_s": "duration_s",
        "zichtbaar_s": "visible_s", "aanwezige_fases": "phases_used"})[[
        "level", "test_id", "observer_phase", "group", "duration_s", "out_of_sight_s",
        "visible_s", "phases_used", "status"]]
    tables = {
        **berekend["per_niveau"],
        "README": pd.DataFrame({"README": uitleg}),
        "details": details,
        "denominators": noemers,
        "warnings": berekend["meldingen"],
        "variables": bron["definities"].drop(columns="kolomnummer"),
        "availability": bron["beschikbaarheid"],
        "phase_selection": selectie,
        "excluded": bron["uitgesloten"],
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)
            ws = writer.book[name]
            # De vijf leidende kolommen van ieder resultaatblad blijven zichtbaar.
            ws.freeze_panes = "F2" if name in NIVEAUS else "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245447")
                ws.column_dimensions[cell.column_letter].width = 24
            if name == "README":
                ws.column_dimensions["A"].width = 140
    return path
