"""Timestamp-Excel row parsing/validation (ticket #94, part of issue #93's
Feature C) — a pure, independently testable module: bytes in, validated
`ExcelRow`s out, no S3/ffmpeg/DB involved (issue #93's "Test seams" /
user story 25).

Only the workbook's first tab is read (CONTEXT.md's Feature C decision);
other tabs are ignored entirely. The expected schema (column headers below)
hasn't been confirmed against a real production sample sheet during this
ticket — same "unconfirmed, revisit" caveat CONTEXT.md already carries for
other external-format assumptions (e.g. app/api/deps.py's JWT email-claim
name) — adjust `_TEST_ID_HEADER`/`_DOG_ID_HEADER`/`_REFERENCE_CAMERA_HEADER`
below if a real sheet uses different header text.
"""

import re
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime, time, timedelta
from io import BytesIO

from openpyxl import load_workbook

_TEST_ID_HEADER = "Test ID"
_DOG_ID_HEADER = "Dog ID"
# The Excel's own reference-camera column (CONTEXT.md: "the Excel's C1/C2
# reference-camera value") — assist only computes a real audio-sync offset
# when both cameras are present, so this tells validation which single
# camera is acceptable when only one was uploaded.
_REFERENCE_CAMERA_HEADER = "C1/C2"

_CONDITIONS = ("ME", "ZE")
_PHASE_NUMBERS = range(1, 9)
# F1-F8 for both conditions — 16 columns, matching the Excel's 16 start-only
# phase timestamps (CONTEXT.md's "ZE_F8 cannot be produced" decision: this
# module parses all 16 faithfully, including ZE_F8's own cell; excluding it
# from a CuttingJob's *expected outputs* is app/services/cutting_jobs.py's
# concern, not this module's).
_PHASE_HEADERS = [f"{condition}_F{n}" for condition in _CONDITIONS for n in _PHASE_NUMBERS]
_REQUIRED_HEADERS = [_TEST_ID_HEADER, _DOG_ID_HEADER, _REFERENCE_CAMERA_HEADER, *_PHASE_HEADERS]

_BARE_NUMERIC_TEST_ID_RE = re.compile(r"^\d+$")
_TEST_ID_RE = re.compile(r"^T\d+$", re.IGNORECASE)
# 3-component (H:MM:SS) and 2-component (M:SS) time strings — a cell that
# fell back to plain text instead of a real Excel time value (e.g. it was
# typed into a text-formatted cell) still parses if it's shaped like either.
_TIME_STRING_HMS_RE = re.compile(r"^(\d{1,2}):(\d{2}):(\d{2})$")
_TIME_STRING_MS_RE = re.compile(r"^(\d{1,2}):(\d{2})$")


class InvalidWorkbookError(Exception):
    """The uploaded bytes aren't a readable .xlsx workbook at all."""


class ExcelSchemaError(Exception):
    """The workbook's first tab is missing one or more required column
    headers (Test ID/Dog ID/C1-C2/F1-F8 ME+ZE)."""

    def __init__(self, missing_headers: list[str]) -> None:
        self.missing_headers = missing_headers
        super().__init__(f"Missing required column(s): {', '.join(missing_headers)}")


class MalformedTimestampCellError(Exception):
    """A phase timestamp cell isn't a real time value — e.g. a literal "-",
    a real example found in an older sample sheet — rather than letting it
    crash `assist`'s own parsing (CONTEXT.md's Feature C decision)."""

    def __init__(self, row_number: int, column: str, raw_value: object) -> None:
        self.row_number = row_number
        self.column = column
        self.raw_value = raw_value
        super().__init__(
            f"Row {row_number}, column {column!r}: not a valid time value ({raw_value!r})"
        )


class MalformedTestIdError(Exception):
    """A row's Test ID cell is blank, or (after bare-numeric normalization)
    still doesn't look like a real Test ID."""

    def __init__(self, row_number: int, raw_value: object) -> None:
        self.row_number = row_number
        self.raw_value = raw_value
        super().__init__(f"Row {row_number}: not a valid Test ID ({raw_value!r})")


class MalformedReferenceCameraError(Exception):
    """A row's reference-camera cell isn't "C1" or "C2"."""

    def __init__(self, row_number: int, raw_value: object) -> None:
        self.row_number = row_number
        self.raw_value = raw_value
        super().__init__(f"Row {row_number}: reference camera must be C1 or C2 ({raw_value!r})")


class TestRowNotFoundError(Exception):
    """No row in the parsed sheet matches the requested Test id (after the
    same bare-numeric normalization applied to the sheet's own Test ID
    cells)."""

    def __init__(self, test_id: str) -> None:
        self.test_id = test_id
        super().__init__(f"No row found for Test {test_id!r}")


@dataclass(frozen=True)
class ExcelRow:
    row_number: int
    test_id: str
    dog_id: str
    # "C1" or "C2".
    reference_camera: str
    # Keyed "ME_F1".."ME_F8"/"ZE_F1".."ZE_F8" -> elapsed seconds from the
    # video's start. A key is present only if its cell wasn't the sheet's
    # 00:00:00-skip marker (CONTEXT.md's Feature C decision) — a phase the
    # Test never ran has no entry here at all, same as a blank cell.
    phase_timestamps: dict[str, int] = field(default_factory=dict)


def normalize_test_id(raw: object) -> str:
    """A bare-numeric Test ID (e.g. "513", seen in an older sample sheet)
    normalized to the `T`-prefixed form used everywhere else in this app
    (CONTEXT.md's Feature C decision); anything else is returned unchanged
    (including already-well-formed and malformed values alike — callers
    validate the result themselves)."""
    text_value = str(raw).strip()
    if _BARE_NUMERIC_TEST_ID_RE.match(text_value):
        return f"T{text_value}"
    return text_value


def _parse_time_cell(value: object, *, row_number: int, column: str) -> int | None:
    """The elapsed seconds `value` (a parsed Excel time value, or a
    plain-text fallback) represents, or None if it's blank or the sheet's
    00:00:00-skip marker (this phase wasn't run).

    Applies the `HH:MM:SS`-means-`MM:SS` conversion: a real, well-known
    Excel time-entry footgun is the likely cause of any 3-component value
    seen here at all — typing e.g. "12:34" (intending 12 minutes 34
    seconds) into a time-formatted cell, Excel parses a 2-component entry
    as H:MM (12 hours, 34 minutes), never M:SS. The two numbers the person
    actually typed land in the hour/minute slots, with seconds always 0 —
    so reading them back as minutes/seconds (`hour * 60 + minute`) recovers
    exactly what was typed, rather than the literal (and nonsensical, for a
    same-day test recording) "12 hours" reading.

    Raises MalformedTimestampCellError for anything that isn't a real time
    value at all (e.g. a stray "-", seen in an older sample sheet).
    """
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None

    # `seconds` is the final elapsed-time value this cell represents;
    # `hour`/`minute`/`second` are only used for the 00:00:00-skip check
    # below, so every branch must set both consistently — in particular the
    # 2-component text branch's `hour` stays 0 (it was never subject to
    # Excel's H:MM misparse, so `seconds` is its minute:second value taken
    # at face value, not the `hour * 60 + minute` reinterpretation the other
    # branches apply), and a 2-component "00:00" must still be caught by
    # the skip check, not treated as "ran at elapsed second 0".
    hour: int
    minute: int
    second: int
    seconds: int
    if isinstance(value, datetime):
        hour, minute, second = value.hour, value.minute, value.second
        seconds = hour * 60 + minute
    elif isinstance(value, time):
        hour, minute, second = value.hour, value.minute, value.second
        seconds = hour * 60 + minute
    elif isinstance(value, timedelta):
        total = int(value.total_seconds())
        hour, minute, second = total // 3600, (total % 3600) // 60, total % 60
        seconds = hour * 60 + minute
    elif isinstance(value, str):
        stripped = value.strip()
        match = _TIME_STRING_HMS_RE.match(stripped)
        if match is not None:
            hour, minute, second = (int(part) for part in match.groups())
            seconds = hour * 60 + minute
        else:
            match = _TIME_STRING_MS_RE.match(stripped)
            if match is None:
                raise MalformedTimestampCellError(row_number, column, value)
            # Already 2-component text, never passed through Excel's own
            # H:MM misparse (it was never a real time-typed cell to begin
            # with) — taken at face value as minutes:seconds.
            hour = 0
            minute, second = (int(part) for part in match.groups())
            seconds = minute * 60 + second
    elif isinstance(value, int | float):
        total = round(value * 86400)
        hour, minute, second = total // 3600, (total % 3600) // 60, total % 60
        seconds = hour * 60 + minute
    else:
        raise MalformedTimestampCellError(row_number, column, value)

    if hour == 0 and minute == 0 and second == 0:
        return None

    return seconds


@contextmanager
def _data_rows(data: bytes) -> Iterator[Iterator[tuple[int, Callable[[str], object]]]]:
    """Open `data`'s first tab, check its header row, and yield an iterator
    of `(row_number, cell)` for every non-blank data row, where
    `cell(header)` reads that row's value in the named column.

    Raises InvalidWorkbookError if `data` isn't a readable .xlsx workbook at
    all, and ExcelSchemaError if the first tab is missing a required column
    header — both workbook-level problems, whichever row a caller wants.
    """
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise InvalidWorkbookError(str(exc)) from exc

    try:
        sheet = workbook.worksheets[0]

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header_index = {
            str(value).strip(): index for index, value in enumerate(header_row) if value is not None
        }
        missing = [header for header in _REQUIRED_HEADERS if header not in header_index]
        if missing:
            raise ExcelSchemaError(missing)

        def rows() -> Iterator[tuple[int, Callable[[str], object]]]:
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if all(value is None for value in row):
                    continue

                def cell(header: str, row: tuple = row) -> object:
                    index = header_index[header]
                    return row[index] if index < len(row) else None

                yield row_number, cell

        yield rows()
    finally:
        workbook.close()


def _parse_row(row_number: int, cell: Callable[[str], object]) -> ExcelRow:
    """One data row, validated in full — raises MalformedTestIdError/
    MalformedReferenceCameraError/MalformedTimestampCellError for its first
    malformed cell (rather than crashing partway, or silently skipping it) —
    CONTEXT.md's Feature C "Ingestion validation" decision."""
    raw_test_id = cell(_TEST_ID_HEADER)
    if raw_test_id is None or not str(raw_test_id).strip():
        raise MalformedTestIdError(row_number, raw_test_id)
    test_id = normalize_test_id(raw_test_id)
    if not _TEST_ID_RE.match(test_id):
        raise MalformedTestIdError(row_number, raw_test_id)
    test_id = test_id.upper()

    raw_dog_id = cell(_DOG_ID_HEADER)
    dog_id = str(raw_dog_id).strip() if raw_dog_id is not None else ""

    raw_camera = cell(_REFERENCE_CAMERA_HEADER)
    reference_camera = str(raw_camera).strip().upper() if raw_camera is not None else ""
    if reference_camera not in ("C1", "C2"):
        raise MalformedReferenceCameraError(row_number, raw_camera)

    phase_timestamps: dict[str, int] = {}
    for header in _PHASE_HEADERS:
        seconds = _parse_time_cell(cell(header), row_number=row_number, column=header)
        if seconds is not None:
            phase_timestamps[header] = seconds

    return ExcelRow(
        row_number=row_number,
        test_id=test_id,
        dog_id=dog_id,
        reference_camera=reference_camera,
        phase_timestamps=phase_timestamps,
    )


def parse_timestamp_workbook(data: bytes) -> list[ExcelRow]:
    """Every data row on `data`'s first tab, validated against the
    Test ID/Dog ID/C1-C2/F1-F8 ME+ZE schema.

    Raises InvalidWorkbookError/ExcelSchemaError for a workbook-level
    problem (see `_data_rows`), and the first malformed cell's own error
    (see `_parse_row`) in any row.
    """
    with _data_rows(data) as rows:
        return [_parse_row(row_number, cell) for row_number, cell in rows]


def read_test_row(data: bytes, test_id: str) -> ExcelRow:
    """The first row on `data`'s first tab whose Test ID matches `test_id`
    (normalized/uppercased on both sides), validated in full — every other
    row is skipped without being validated at all.

    Ticket #97: a batch submission shares one workbook across up to five
    Tests, so another row's malformed cell (or a row for a Test not being
    cut at all) must never fail this Test. A row whose own Test ID cell is
    blank or malformed can't be attributed to any Test, so it's skipped too
    — if it was meant to be `test_id`'s, this raises TestRowNotFoundError.
    Workbook-level problems still raise InvalidWorkbookError/
    ExcelSchemaError, as they would for any Test.
    """
    wanted = normalize_test_id(test_id).upper()
    with _data_rows(data) as rows:
        for row_number, cell in rows:
            if normalize_test_id(cell(_TEST_ID_HEADER)).upper() == wanted:
                return _parse_row(row_number, cell)
    raise TestRowNotFoundError(test_id)


def find_test_row(rows: list[ExcelRow], test_id: str) -> ExcelRow:
    """The row matching `test_id` (normalized/uppercased the same way the
    sheet's own Test ID cells are), or raise TestRowNotFoundError."""
    normalized = normalize_test_id(test_id).upper()
    for row in rows:
        if row.test_id == normalized:
            return row
    raise TestRowNotFoundError(test_id)
