"""Build synthetic_observer_export.xlsx from a real Observer export's *shape*.

Copies only structure from the real export — the header row of its raw
`Results` sheet, the only sheet consolidation reads (ticket #151) — and
fills every data cell with generated values for two fake tests (T901/T902,
dogs D901/D902) across all 15 phases. No measurement, test ID, dog ID or
other metadata value from the real export is copied, and none of the
hand-made working copies (`Results (2)`, `Results (REL)`) is written.

The values are generated to satisfy consolidation/'s own validation: every
metric is numeric or `-` (Observer's measured zero), counts are whole
numbers, and out-of-sight time and each behaviour's duration fit inside the
phase.

Re-run from backend/ (only needed if the real export's column layout
changes):

    PYTHONPATH=.. .venv/bin/python \\
        tests/fixtures/consolidation/build_synthetic_observer_export.py \\
        /path/to/real_observer_export.xlsx
"""

import random
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

OUTPUT = Path(__file__).with_name("synthetic_observer_export.xlsx")
TESTS = [("T901", "D901", "M"), ("T902", "D902", "V")]
PHASE_DURATION_S = 120.0
OOS_DURATION_S = 10.0


def _value(header: str, rng: random.Random) -> float | int | str:
    if header == "Duration":
        return PHASE_DURATION_S
    if header.startswith("Total duration Out of sight"):
        value = OOS_DURATION_S
    elif header.startswith("Total duration "):
        value = rng.choice([0.0, 0.0, 0.0, 4.5, 12.25, 30.0])
    else:
        value = rng.choice([0, 0, 1, 2, 3])
    # Real exports write "-" for many zero measurements.
    return "-" if value == 0 and rng.random() < 0.5 else value


def _cell(header: str, structural: dict[str, str], rng: random.Random) -> float | int | str:
    if header in structural:
        return structural[header]
    if header == "Duration" or header.startswith("Total "):
        return _value(header, rng)
    # Observer's own container columns are never read.
    return "synthetic"


def build(real_export: Path) -> None:
    rng = random.Random(151)
    real = load_workbook(real_export, read_only=True)
    headers = [str(h) for h in next(real["Results"].iter_rows(max_row=1, values_only=True))]
    real.close()

    workbook = Workbook()
    results = workbook.active
    results.title = "Results"
    results.append(headers)
    for test_id, dog_id, sex in TESTS:
        for phase in range(1, 16):
            structural = {
                "Observations": f"{test_id}_synthetic_F{phase}",
                "Aanwezigheid FP in fase": "True" if phase <= 7 else "False",
                "Dog ID": dog_id,
                "Fase": str(phase if phase <= 8 else phase - 8),
                "Geslacht hond": sex,
                "Test ID": test_id,
            }
            results.append([_cell(h, structural, rng) for h in headers])
    workbook.save(OUTPUT)


if __name__ == "__main__":
    build(Path(sys.argv[1]))
    print(f"Wrote {OUTPUT}")
