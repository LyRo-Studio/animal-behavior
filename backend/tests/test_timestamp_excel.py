"""Direct tests for app/services/timestamp_excel.py (ticket #94, part of
issue #93's Feature C) — pure parsing/validation, no DB/S3/ffmpeg (issue
#93's Test seams / user story 25). Fixture workbooks are built on the fly
with openpyxl rather than checked in as binary files, so each test's exact
cell content is visible right next to its assertions.
"""

from datetime import time
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.services.timestamp_excel import (
    ExcelSchemaError,
    InvalidWorkbookError,
    MalformedReferenceCameraError,
    MalformedTestIdError,
    MalformedTimestampCellError,
    TestRowNotFoundError,
    find_test_row,
    normalize_test_id,
    parse_timestamp_workbook,
)

_CONDITIONS = ("ME", "ZE")
_PHASE_HEADERS = [f"{condition}_F{n}" for condition in _CONDITIONS for n in range(1, 9)]
_HEADERS = ["Test ID", "Dog ID", "C1/C2", *_PHASE_HEADERS]


def _workbook_bytes(rows: list[list[object]], *, headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers if headers is not None else _HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _row(test_id, dog_id="Rex", reference_camera="C1", **phase_overrides) -> list[object]:
    """A full data row, every phase defaulting to blank (skipped) unless
    overridden by keyword (e.g. `ME_F1=time(0, 5, 23)`)."""
    values = [test_id, dog_id, reference_camera]
    values += [phase_overrides.get(header) for header in _PHASE_HEADERS]
    return values


def test_parses_a_well_formed_row():
    """`time(5, 23, 0)`: the sheet's own convention for "5 minutes 23
    seconds" — a 2-component entry Excel always parses as H:MM (never
    M:SS), landing the typed 5 and 23 in the hour/minute slots with seconds
    always 0 (see `_parse_time_cell`'s docstring)."""
    data = _workbook_bytes([_row("T001", reference_camera="C2", ME_F1=time(5, 23, 0))])

    rows = parse_timestamp_workbook(data)

    assert len(rows) == 1
    row = rows[0]
    assert row.test_id == "T001"
    assert row.dog_id == "Rex"
    assert row.reference_camera == "C2"
    assert row.phase_timestamps == {"ME_F1": 5 * 60 + 23}


def test_only_the_first_tab_is_read():
    data = _workbook_bytes([_row("T001")])
    workbook_bytes = BytesIO(data)
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_bytes)
    extra = workbook.create_sheet("second tab")
    extra.append(["not", "the", "expected", "schema", "at", "all"])
    buffer = BytesIO()
    workbook.save(buffer)

    rows = parse_timestamp_workbook(buffer.getvalue())

    assert [row.test_id for row in rows] == ["T001"]


def test_bare_numeric_test_id_is_normalized():
    data = _workbook_bytes([_row(513)])

    rows = parse_timestamp_workbook(data)

    assert rows[0].test_id == "T513"


def test_normalize_test_id_helper_matches_workbook_behavior():
    assert normalize_test_id("513") == "T513"
    assert normalize_test_id(513) == "T513"
    assert normalize_test_id("T041") == "T041"


def test_malformed_test_id_raises_with_row_number():
    data = _workbook_bytes([_row("not-a-test-id")])

    with pytest.raises(MalformedTestIdError) as exc_info:
        parse_timestamp_workbook(data)
    assert exc_info.value.row_number == 2


def test_blank_test_id_raises():
    data = _workbook_bytes([_row(None)])

    with pytest.raises(MalformedTestIdError):
        parse_timestamp_workbook(data)


def test_reference_camera_is_case_normalized():
    data = _workbook_bytes([_row("T001", reference_camera="c1")])

    rows = parse_timestamp_workbook(data)

    assert rows[0].reference_camera == "C1"


def test_malformed_reference_camera_raises():
    data = _workbook_bytes([_row("T001", reference_camera="C3")])

    with pytest.raises(MalformedReferenceCameraError) as exc_info:
        parse_timestamp_workbook(data)
    assert exc_info.value.row_number == 2


def test_missing_header_raises_schema_error():
    headers = [h for h in _HEADERS if h != "Dog ID"]
    data = _workbook_bytes([["T001", "C1"]], headers=headers)

    with pytest.raises(ExcelSchemaError) as exc_info:
        parse_timestamp_workbook(data)
    assert "Dog ID" in exc_info.value.missing_headers


def test_not_a_workbook_raises_invalid_workbook_error():
    with pytest.raises(InvalidWorkbookError):
        parse_timestamp_workbook(b"this is not an xlsx file")


def test_00_00_00_is_skipped_not_recorded():
    data = _workbook_bytes([_row("T001", ME_F1=time(0, 0, 0))])

    rows = parse_timestamp_workbook(data)

    assert "ME_F1" not in rows[0].phase_timestamps


def test_blank_phase_cell_is_skipped():
    data = _workbook_bytes([_row("T001")])

    rows = parse_timestamp_workbook(data)

    assert rows[0].phase_timestamps == {}


def test_hh_mm_ss_shaped_cell_is_read_as_mm_ss():
    """A 12:34:00 time value (a person typing "12:34" into a time-formatted
    cell, which Excel parses as 12 hours 34 minutes rather than 12 minutes
    34 seconds) means 12 minutes 34 seconds, not 12 hours 34 minutes."""
    data = _workbook_bytes([_row("T001", ME_F1=time(12, 34, 0))])

    rows = parse_timestamp_workbook(data)

    assert rows[0].phase_timestamps["ME_F1"] == 12 * 60 + 34


def test_2_component_00_00_text_is_skipped_not_recorded_as_zero():
    """A 2-component "00:00" cell (plain text, never passed through Excel's
    own time-type parsing) is the skip marker just like "00:00:00" is —
    must not be treated as "ran at elapsed second 0"."""
    data = _workbook_bytes([_row("T001", ME_F1="00:00")])

    rows = parse_timestamp_workbook(data)

    assert "ME_F1" not in rows[0].phase_timestamps


def test_2_component_text_cell_is_read_as_mm_ss_directly():
    """Unlike a genuine (Excel-typed) 3-component value, 2-component text
    was never subject to Excel's H:MM misparse — taken at face value."""
    data = _workbook_bytes([_row("T001", ME_F1="12:34")])

    rows = parse_timestamp_workbook(data)

    assert rows[0].phase_timestamps["ME_F1"] == 12 * 60 + 34


def test_malformed_timestamp_cell_raises_with_row_and_column():
    data = _workbook_bytes([_row("T001", ME_F3="-")])

    with pytest.raises(MalformedTimestampCellError) as exc_info:
        parse_timestamp_workbook(data)
    assert exc_info.value.row_number == 2
    assert exc_info.value.column == "ME_F3"


def test_malformed_cell_does_not_crash_the_whole_parse_silently():
    """A malformed cell raises a specific error rather than being silently
    skipped or crashing with an unrelated exception."""
    data = _workbook_bytes([_row("T001", ZE_F8="not a time")])

    with pytest.raises(MalformedTimestampCellError):
        parse_timestamp_workbook(data)


def test_ze_f8_is_parsed_like_any_other_phase():
    """ZE_F8 can never become an expected CuttingJobOutput (structurally
    impossible — CONTEXT.md), but this module still parses its cell
    faithfully; excluding it is cutting_jobs.py's concern."""
    data = _workbook_bytes([_row("T001", ZE_F8=time(10, 0, 0))])

    rows = parse_timestamp_workbook(data)

    assert rows[0].phase_timestamps["ZE_F8"] == 600


def test_multiple_rows_are_all_parsed():
    data = _workbook_bytes([_row("T001"), _row("T002", reference_camera="C2")])

    rows = parse_timestamp_workbook(data)

    assert [row.test_id for row in rows] == ["T001", "T002"]


def test_trailing_blank_row_is_ignored():
    data = _workbook_bytes([_row("T001"), [None] * len(_HEADERS)])

    rows = parse_timestamp_workbook(data)

    assert len(rows) == 1


def test_find_test_row_returns_the_matching_row():
    data = _workbook_bytes([_row("T001"), _row("T002")])
    rows = parse_timestamp_workbook(data)

    found = find_test_row(rows, "T002")

    assert found.test_id == "T002"


def test_find_test_row_normalizes_a_bare_numeric_lookup():
    data = _workbook_bytes([_row(513)])
    rows = parse_timestamp_workbook(data)

    found = find_test_row(rows, "513")

    assert found.test_id == "T513"


def test_find_test_row_raises_for_no_match():
    data = _workbook_bytes([_row("T001")])
    rows = parse_timestamp_workbook(data)

    with pytest.raises(TestRowNotFoundError):
        find_test_row(rows, "T404")
